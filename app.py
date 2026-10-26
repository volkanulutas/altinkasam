import csv
import html
import io
import math
import os
import re
import time
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta

import requests
from flask import Flask, jsonify, request, send_from_directory, session
from flask_sqlalchemy import SQLAlchemy
from openpyxl import load_workbook
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

URL = "https://www.sedatascikuyumculuk.com/currencies"
CACHE_SECONDS = 60
NUM = r"(\d{1,3}(?:\.\d{3})*,\d{2})"

# key, ekranda görünen ad, birim, sitedeki ad
URUNLER = [
    ("gram", "Gram Altın (24 ayar)", "gram", "Gram Altın 1gr 24 Ayar"),
    ("bilezik", "Bilezik 22 Ayar", "gram", "Bilezik 22 Ayar"),
    ("ceyrek", "Çeyrek Altın", "adet", "Yeni Tarih Çeyrek"),
    ("ata", "Ata Lira", "adet", "Ata Lira"),
]
URUN_KEYS = {u[0] for u in URUNLER}
MAX_ALIM = 5000

# ---------------------------------------------------------------- ayarlar
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__, static_folder=None)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)

on_render = bool(os.environ.get("RENDER"))
secret = os.environ.get("SECRET_KEY")
if not secret:
    if on_render:
        raise RuntimeError("SECRET_KEY ortam değişkeni tanımlı değil.")
    secret = "yerel-gelistirme-anahtari"

db_url = os.environ.get("DATABASE_URL", "sqlite:///altinkasam.db")
if db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)

app.config.update(
    SECRET_KEY=secret,
    SQLALCHEMY_DATABASE_URI=db_url,
    SQLALCHEMY_ENGINE_OPTIONS={"pool_pre_ping": True},
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=on_render,
    PERMANENT_SESSION_LIFETIME=timedelta(days=30),
    MAX_CONTENT_LENGTH=4 * 1024 * 1024,
)
db = SQLAlchemy(app)


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(30), unique=True, nullable=False, index=True)
    password_hash = db.Column(db.String(300), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Purchase(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    tarih = db.Column(db.Date, nullable=False)
    urun = db.Column(db.String(20), nullable=False)
    miktar = db.Column(db.Float, nullable=False)
    maliyet = db.Column(db.Float, nullable=False)
    vadeli = db.Column(db.Boolean, nullable=False, default=False)

    def to_dict(self):
        return {"id": self.id, "tarih": self.tarih.isoformat(), "urun": self.urun,
                "miktar": self.miktar, "maliyet": self.maliyet, "vadeli": self.vadeli}


with app.app_context():
    db.create_all()

# ---------------------------------------------------------------- güvenlik
_fails = defaultdict(list)


def throttled(key, limit, window):
    now = time.time()
    _fails[key] = [t for t in _fails[key] if now - t < window]
    return len(_fails[key]) >= limit


def record(key):
    _fails[key].append(time.time())


@app.before_request
def csrf_ve_giris_kontrolu():
    # Durum değiştiren isteklerde özel başlık zorunlu (başka sitelerden form ile gönderilemez)
    if request.method in ("POST", "PUT", "DELETE") and request.path.startswith("/api/"):
        if request.headers.get("X-Requested-With") != "fetch":
            return jsonify({"hata": "Geçersiz istek."}), 400


@app.after_request
def baslik(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    if request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


def giris_gerekli(fn):
    from functools import wraps

    @wraps(fn)
    def w(*a, **k):
        if not session.get("uid"):
            return jsonify({"hata": "Giriş yapmalısınız."}), 401
        return fn(*a, **k)
    return w


USERNAME_RE = re.compile(r"^[a-z0-9_.-]{3,30}$")
_DUMMY_HASH = generate_password_hash("dummy-password")


def kimlik_bilgisi():
    d = request.get_json(silent=True) or {}
    return str(d.get("kullanici", "")).strip().lower(), str(d.get("sifre", ""))


@app.post("/api/kayit")
def kayit():
    ip = request.remote_addr or "?"
    if throttled("reg:" + ip, 10, 3600):
        return jsonify({"hata": "Çok fazla deneme. Daha sonra tekrar deneyin."}), 429
    record("reg:" + ip)
    u, p = kimlik_bilgisi()
    if not USERNAME_RE.match(u):
        return jsonify({"hata": "Kullanıcı adı 3-30 karakter olmalı (harf, rakam, . _ -)."}), 400
    if len(p) < 8 or len(p) > 200:
        return jsonify({"hata": "Şifre en az 8 karakter olmalı."}), 400
    if User.query.filter_by(username=u).first():
        return jsonify({"hata": "Bu kullanıcı adı alınmış."}), 409
    user = User(username=u, password_hash=generate_password_hash(p))
    db.session.add(user)
    db.session.commit()
    session.clear()
    session.permanent = True
    session["uid"] = user.id
    return jsonify({"kullanici": user.username})


@app.post("/api/giris")
def giris():
    ip = request.remote_addr or "?"
    u, p = kimlik_bilgisi()
    if throttled("ip:" + ip, 20, 600) or throttled("u:" + u, 10, 600):
        return jsonify({"hata": "Çok fazla hatalı deneme. 10 dakika sonra tekrar deneyin."}), 429
    user = User.query.filter_by(username=u).first()
    ok = check_password_hash(user.password_hash if user else _DUMMY_HASH, p)
    if not (user and ok):
        record("ip:" + ip)
        record("u:" + u)
        return jsonify({"hata": "Kullanıcı adı veya şifre hatalı."}), 401
    session.clear()
    session.permanent = True
    session["uid"] = user.id
    return jsonify({"kullanici": user.username})


@app.post("/api/cikis")
def cikis():
    session.clear()
    return jsonify({"ok": True})


@app.get("/api/ben")
def ben():
    uid = session.get("uid")
    user = db.session.get(User, uid) if uid else None
    if not user:
        session.clear()
        return jsonify({"hata": "Giriş yapılmamış."}), 401
    return jsonify({"kullanici": user.username})


# ---------------------------------------------------------------- alımlar
def bugun():
    return date.today()


def alim_ayikla(d):
    if not isinstance(d, dict):
        raise ValueError("Geçersiz kayıt.")
    urun = str(d.get("urun", ""))
    if urun not in URUN_KEYS:
        raise ValueError("Geçersiz ürün.")
    try:
        tarih = date.fromisoformat(str(d.get("tarih", ""))[:10])
        miktar = float(d.get("miktar"))
        maliyet = float(d.get("maliyet"))
    except (TypeError, ValueError):
        raise ValueError("Geçersiz tarih veya sayı.")
    if tarih < date(1990, 1, 1) or tarih > bugun() + timedelta(days=1):
        raise ValueError("Tarih geçersiz.")
    if not (math.isfinite(miktar) and 0 < miktar <= 1_000_000):
        raise ValueError("Miktar geçersiz.")
    if not (math.isfinite(maliyet) and 0 <= maliyet <= 1_000_000_000):
        raise ValueError("Birim maliyet geçersiz.")
    return dict(tarih=tarih, urun=urun, miktar=miktar, maliyet=maliyet, vadeli=bool(d.get("vadeli")))


def alim_listesi(uid):
    q = Purchase.query.filter_by(user_id=uid).order_by(Purchase.tarih.desc(), Purchase.id.desc())
    return [p.to_dict() for p in q]


@app.get("/api/alimlar")
@giris_gerekli
def alimlar_getir():
    return jsonify(alim_listesi(session["uid"]))


@app.post("/api/alimlar")
@giris_gerekli
def alim_ekle():
    uid = session["uid"]
    if Purchase.query.filter_by(user_id=uid).count() >= MAX_ALIM:
        return jsonify({"hata": "Kayıt sınırına ulaşıldı."}), 400
    try:
        v = alim_ayikla(request.get_json(silent=True))
    except ValueError as e:
        return jsonify({"hata": str(e)}), 400
    p = Purchase(user_id=uid, **v)
    db.session.add(p)
    db.session.commit()
    return jsonify(p.to_dict()), 201


@app.post("/api/alimlar/toplu")
@giris_gerekli
def alim_toplu():
    uid = session["uid"]
    liste = (request.get_json(silent=True) or {}).get("alimlar")
    if not isinstance(liste, list) or len(liste) > 1000:
        return jsonify({"hata": "Geçersiz veri (en fazla 1000 kayıt)."}), 400
    mevcut = Purchase.query.filter_by(user_id=uid).count()
    eklenen = 0
    for d in liste:
        if mevcut + eklenen >= MAX_ALIM:
            break
        try:
            db.session.add(Purchase(user_id=uid, **alim_ayikla(d)))
            eklenen += 1
        except ValueError:
            continue
    db.session.commit()
    return jsonify({"eklenen": eklenen, "alimlar": alim_listesi(uid)})


# ---------------------------------------------------------------- Excel / CSV içe aktarma
def tr_norm(s):
    s = str(s if s is not None else "").strip().replace("İ", "i").replace("I", "i").replace("ı", "i")
    s = unicodedata.normalize("NFKD", s.lower())
    return "".join(c for c in s if not unicodedata.combining(c))


def kolon_bul(satir):
    m = {}
    for i, h in enumerate(satir):
        n = tr_norm(h)
        for anahtar, onek in (("tur", "tur"), ("adet", "adet"), ("tarih", "tarih"),
                              ("alis", "alis"), ("vadeli", "vadeli")):
            if n.startswith(onek) and anahtar not in m:
                m[anahtar] = i
    return m if {"tur", "adet", "tarih", "alis"} <= m.keys() else None


def urun_bul(t):
    n = tr_norm(t)
    if "bilezik" in n:
        return "bilezik"
    if "ceyrek" in n:
        return "ceyrek"
    if n == "ata" or n.startswith("ata "):
        return "ata"
    if "gram" in n:
        return "gram"
    return None


def sayi(v):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("₺", "").replace("TL", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return float(s)


def tarih_oku(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for f in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    raise ValueError("tarih okunamadı")


def tablo_satirlari(ad, veri):
    """(sayfa_adı, satırlar) döndürür; Excel'de 'ALTIN' sekmesi öncelikli."""
    if ad.lower().endswith(".csv"):
        try:
            metin = veri.decode("utf-8-sig")
        except UnicodeDecodeError:
            metin = veri.decode("cp1254", errors="replace")
        ilk = metin.splitlines()[0] if metin else ""
        ayrac = ";" if ilk.count(";") >= ilk.count(",") else ","
        return [("CSV", list(csv.reader(io.StringIO(metin), delimiter=ayrac)))]
    wb = load_workbook(io.BytesIO(veri), read_only=True, data_only=True)
    sheets = sorted(wb.worksheets, key=lambda w: tr_norm(w.title) != "altin")
    return [(w.title, [list(r) for r in w.iter_rows(values_only=True, max_row=3000)]) for w in sheets]


@app.post("/api/ice-aktar")
@giris_gerekli
def ice_aktar():
    f = request.files.get("dosya")
    ad = (f.filename or "").lower() if f else ""
    if not ad.endswith((".xlsx", ".csv")):
        return jsonify({"hata": "Lütfen .xlsx veya .csv dosyası seçin. (Google Sheets: Dosya → İndir → Microsoft Excel)"}), 400
    try:
        sayfalar = tablo_satirlari(ad, f.read())
    except Exception:
        return jsonify({"hata": "Dosya okunamadı."}), 400
    for sayfa, rows in sayfalar:
        for hi, row in enumerate(rows[:15]):
            kol = kolon_bul(row)
            if not kol:
                continue
            satirlar, hatalar = [], []
            for ri in range(hi + 1, len(rows)):
                r = rows[ri]
                al = lambda k: r[kol[k]] if k in kol and kol[k] < len(r) else None
                if str(al("tur") or "").strip() == "" and str(al("adet") or "").strip() == "":
                    continue
                no = ri + 1
                urun = urun_bul(al("tur"))
                if not urun:
                    hatalar.append(f"Satır {no}: tanınmayan tür '{al('tur')}'")
                    continue
                try:
                    miktar, fiyat, tarih = sayi(al("adet")), sayi(al("alis")), tarih_oku(al("tarih"))
                    if not miktar or miktar <= 0 or fiyat is None or fiyat < 0:
                        raise ValueError
                except (ValueError, TypeError):
                    hatalar.append(f"Satır {no}: adet, tarih veya fiyat okunamadı")
                    continue
                satirlar.append({"satir": no, "tarih": tarih.isoformat(), "urun": urun, "miktar": miktar,
                                 "maliyet": fiyat, "vadeli": tr_norm(al("vadeli")) in ("evet", "e", "true", "1", "x")})
            if len(satirlar) > 1000:
                return jsonify({"hata": "En fazla 1000 satır içe aktarılabilir."}), 400
            return jsonify({"sayfa": sayfa, "satirlar": satirlar, "hatalar": hatalar})
    return jsonify({"hata": "Başlık satırı bulunamadı. Sütunlar: TÜR, ADET, TARİH, ALIŞ FİYATI, VADELİ olmalı."}), 400


@app.errorhandler(413)
def cok_buyuk(_):
    return jsonify({"hata": "Dosya çok büyük (en fazla 4 MB)."}), 413


@app.delete("/api/alimlar/<int:pid>")
@giris_gerekli
def alim_sil(pid):
    p = Purchase.query.filter_by(id=pid, user_id=session["uid"]).first()
    if not p:
        return jsonify({"hata": "Kayıt bulunamadı."}), 404
    db.session.delete(p)
    db.session.commit()
    return jsonify({"ok": True})


# ---------------------------------------------------------------- fiyatlar
_cache = {"t": 0, "data": None}


def to_float(s):
    return float(s.replace(".", "").replace(",", "."))


def sayfa_metni():
    r = requests.get(URL, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    r.encoding = "utf-8"
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", r.text)
    t = re.sub(r"<[^>]+>", " ", t)
    t = html.unescape(t)
    return re.sub(r"\s+", " ", t)


def parse(metin):
    m = re.search(r"Son güncelleme:?\s*(\d{1,2} \w+ \d{4} \d{2}:\d{2})", metin)
    fiyatlar = []
    for key, ad, birim, site in URUNLER:
        f = re.search(re.escape(site) + r"\s*" + NUM + r"\s*" + NUM, metin)
        if not f:
            raise ValueError(f"Sitede bulunamadı: {site}")
        fiyatlar.append({"key": key, "ad": ad, "birim": birim,
                         "alis": to_float(f.group(1)), "satis": to_float(f.group(2))})
    return {"kaynak_zaman": m.group(1) if m else None, "fiyatlar": fiyatlar}


@app.get("/api/fiyatlar")
def fiyatlar():
    now = time.time()
    if _cache["data"] and now - _cache["t"] < CACHE_SECONDS:
        return jsonify({**_cache["data"], "eski": False})
    try:
        data = parse(sayfa_metni())
        _cache.update(t=now, data=data)
        return jsonify({**data, "eski": False})
    except Exception as e:
        if _cache["data"]:
            return jsonify({**_cache["data"], "eski": True})
        return jsonify({"hata": str(e)}), 502


@app.route("/")
def index():
    return send_from_directory(BASE_DIR, "index.html")


if __name__ == "__main__":
    app.run(debug=True)
